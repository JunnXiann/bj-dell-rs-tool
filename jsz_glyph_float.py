from pymongo import MongoClient
import helper as hp
import pandas as pd
from openpyxl import load_workbook
from openpyxl.drawing.image import Image as OpenpyxlImage
from openpyxl.utils import get_column_letter
import sys
import os

db = hp.get_db('tw-aux')  # uses mongodb-tw-aux in app.yml
coll = db['glyph_jsz']

BATCH = "L2-2R1"
OUT_XLSX = "glyph_jsz_L2-2R1.xlsx"
BASE_DES = "/nas/web-static/tw-aux/des_imgs"
BASE_HANYI = "/nas/web-static/tw-aux/hanyi_chars"
BASE_REV1 = "/nas/web-static/tw-aux/review1_imgs"
BASE_REV2 = "/nas/web-static/tw-aux/review2_imgs"

def char_image_path(char_name):
    if not char_name:
        return ""
    parts = char_name.split('_')
    if len(parts) >= 2:
        prefix = parts[0]
        rest = parts[1:3]
        return f"{BASE_HANYI}/{prefix}/" + "/".join(rest) + f"/{char_name}.jpg"
    return f"{BASE_HANYI}/{char_name}.jpg"

def _validate_base_dirs(paths):
    bad = []
    for p in paths:
        if not p:
            bad.append((p, "empty path"))
            continue
        if not os.path.isabs(p):
            p = os.path.normpath(os.path.join(os.getcwd(), p))
        if not os.path.isdir(p):
            bad.append((p, "not a directory"))
            continue
        if not os.access(p, os.R_OK):
            bad.append((p, "no read permission"))
    if bad:
        print("ERROR: required base directories are not accessible:")
        for p, reason in bad:
            print(f" - {p}: {reason}")
        sys.exit(2)

_validate_base_dirs([BASE_DES, BASE_HANYI, BASE_REV1, BASE_REV2])


query = {"review_work_batch": {"$regex": f'^{BATCH}'}}
proj = {
    "unicode": 1, "std": 1, "txt": 1, "description": 1, "char_name": 1,
    "txt_kind": 1, "review1": 1, "review2": 1, "review3": 1, "pass1": 1, "review6": 1,
    "std_cnt": 1, "des_img": 1, "review1_img": 1, "review2_img": 1
}

cursor = coll.find(query, proj).sort([("std_cnt", -1), ("std", 1)])
rows = []
for d in cursor:
    uni = d.get("unicode")
    # only include image paths for these three when corresponding flag == '是'
    des_path = f"{BASE_DES}/{uni}.jpg" if uni and d.get("des_img") == "是" else ""
    rev1_path = f"{BASE_REV1}/{uni}.jpg" if uni and d.get("review1_img") == "是" else ""
    rev2_path = f"{BASE_REV2}/{uni}.jpg" if uni and d.get("review2_img") == "是" else ""
    rows.append({
        "unicode": uni,
        "正字": d.get("std", ""),
        "字種": d.get("txt", ""),
        "字體特徵說明": d.get("description", ""),
        "字體特徵說明圖片": des_path,
        "當前字圖圖片": char_image_path(d.get("char_name", "")),
        "當前字圖編碼": d.get("char_name", ""),
        "計算機字（字体效果）": d.get("txt_kind", ""),
        "审1意見": d.get("review1", ""),
        "审2意見": d.get("review2", ""),
        "审3意見": d.get("review3", ""),
        "截圖1": rev1_path,
        "截圖2": rev2_path,
        "通過1": d.get("pass1", ""),
        "綜合意見1": d.get("review6", "")
    })

df = pd.DataFrame(rows)
df.to_excel(OUT_XLSX, index=False)

# insert images without scaling; ensure image anchored inside cell by adjusting row height and column width
IMG_COLS = ["字體特徵說明圖片", "當前字圖圖片", "截圖1", "截圖2"]
wb = load_workbook(OUT_XLSX)
ws = wb.active
headers = {cell.value: idx for idx, cell in enumerate(ws[1], start=1)}

for row_idx in range(2, ws.max_row + 1):
    for name in IMG_COLS:
        col_idx = headers.get(name)
        if not col_idx:
            continue
        cell = ws.cell(row=row_idx, column=col_idx)
        img_path = cell.value
        if not img_path:
            continue
        if not os.path.isabs(img_path):
            img_path = os.path.normpath(os.path.join(os.getcwd(), img_path))
        if not os.path.exists(img_path):
            continue
        if not os.path.exists(img_path):
            print(f"⚠️ Missing file skipped: {img_path}")
            continue
            
        try:
            img = OpenpyxlImage(img_path)  # native size, do not scale
            
            # adjust row height to fit image (points ≈ pixels * 0.75)
            needed_row_h = int(img.height * 0.75) or None
            if needed_row_h:
                cur_h = ws.row_dimensions[row_idx].height
                if not cur_h or needed_row_h > cur_h:
                    ws.row_dimensions[row_idx].height = needed_row_h
                    
            # adjust column width to fit image (approx pixels -> excel width: (pixels - 5)/7)
            col_letter = get_column_letter(col_idx)
            needed_col_w = max(1, (img.width - 5) / 7)
            cur_w = ws.column_dimensions[col_letter].width
            if not cur_w or needed_col_w > cur_w:
                ws.column_dimensions[col_letter].width = needed_col_w
                
            # clear cell text and add image anchored to cell
            cell.value = None
            anchor = f"{col_letter}{row_idx}"
            ws.add_image(img, anchor)
            
        except ImportError:
            print("🚨 ERROR: The 'Pillow' library is required to insert images.")
            print("Please run: pip install Pillow")
            sys.exit(1)
        except Exception as e:
            # Now it will actually tell you WHY it failed!
            print(f"❌ Error inserting {img_path}: {e}")
            continue

wb.save(OUT_XLSX)
print(f"Exported {len(df)} rows to {OUT_XLSX} (images embedded where files found)")