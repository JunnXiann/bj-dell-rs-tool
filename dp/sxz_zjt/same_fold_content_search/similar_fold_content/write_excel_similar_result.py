# -*- coding: utf-8 -*-
"""
@File    : write_excel_similar_result.py
@Time    : 2025/5/27 15:29
@Author  : zhujiantao
@Version : 1.0
@Desc    : 将内容相似扣信息写入到excel中，包含在excel中嵌入扣图
注意：similar_fold中扣图只保留2张样例
"""
import json
import os
import uuid
from openpyxl import Workbook
from openpyxl.styles import Alignment
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import get_column_letter
from PIL import Image


def auto_adjust_column_width(ws):
    """自动调整非图像列的列宽"""
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        if col_letter == 'F':  # 图片列不设置宽度
            continue
        max_length = 0
        for cell in col:
            try:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            except:
                pass
        ws.column_dimensions[col_letter].width = max_length + 2


def json_to_excel_with_embedded_images(data, output_file):
    wb = Workbook()
    ws = wb.active
    ws.title = "Similar Folds with Images"

    headers = [
        "fold_content", "same_han", "same_volume",
        "fold_url", "pdf_page_url", "fold_image",
        "similar_percentage", "same_pdf_page", "adjacent"
    ]
    ws.append(headers)

    row = 2
    for item in data:
        fold_content = item.get("fold_content", "")
        same_han = item.get("same_han", "")
        same_volume = item.get("same_volume", "")
        similar_percentage = item.get("similar_percentage", "")
        same_pdf_page = item.get("same_pdf_page", "")
        adjacent = item.get("adjacent", "")
        id_list = item.get("fold_id_list", [])

        start_row = row
        for sub in id_list:
            ws.cell(row=row, column=4, value=sub.get("fold_url", ""))
            ws.cell(row=row, column=5, value=sub.get("pdf_page_url", ""))

            image_path = sub.get("fold_image_path", "")
            if os.path.exists(image_path):
                try:
                    im = Image.open(image_path)
                    w, h = im.size
                    im = im.resize((w // 2, h // 2))

                    # 生成唯一临时图像路径
                    temp_path = f"temp/~temp_{uuid.uuid4().hex}.jpg"
                    im.save(temp_path)

                    img = XLImage(temp_path)
                    img.anchor = f"F{row}"
                    ws.add_image(img)

                    # 设置行高/列宽匹配图像尺寸
                    ws.row_dimensions[row].height = (h // 2) * 0.75
                    ws.column_dimensions["F"].width = (w // 2) / 7

                except Exception as e:
                    print(f"⚠️ 图片处理失败: {image_path} - {e}")

            row += 1

        if id_list:
            end_row = row - 1
            for col, val in zip([1, 2, 3, 7, 8, 9],
                                [fold_content, same_han, same_volume, similar_percentage, same_pdf_page, adjacent]):
                ws.merge_cells(start_row=start_row, start_column=col, end_row=end_row, end_column=col)
                cell = ws.cell(row=start_row, column=col, value=val)
                cell.alignment = Alignment(vertical="center", horizontal="center")

    auto_adjust_column_width(ws)
    wb.save(output_file)
    print(f"✅ Excel 成功保存为：{output_file}")


# 示例：从 JSON 文件读取数据
with open("updated_similar_results.json", "r", encoding="utf-8") as f:
    json_data = json.load(f)

json_to_excel_with_embedded_images(json_data, "相似折图结果_自动列宽.xlsx")
