import json

from openpyxl import load_workbook

def read_excel_to_dict(excel_path: str):
    wb = load_workbook(excel_path)
    ws = wb.active

    result = {}
    for row in ws.iter_rows(min_row=2, values_only=True):  # 跳过标题行
        key = str(row[0]).strip()
        pdf_page = row[1]
        pdf_page_url = str(row[2]).strip()

        result[key] = {
            "pdf_page": pdf_page,
            "pdf_page_url": pdf_page_url
        }

    return result

# 示例调用
excel_path = "思溪藏扣图对应PDF页图统计表-20250415 .xlsx"  # 替换为实际文件路径
data_dict = read_excel_to_dict(excel_path)

# 打印结果预览
from pprint import pprint
pprint(data_dict)

with open("fold_pdf_pge_url.json", "w") as f:
    f.write(json.dumps(data_dict))





