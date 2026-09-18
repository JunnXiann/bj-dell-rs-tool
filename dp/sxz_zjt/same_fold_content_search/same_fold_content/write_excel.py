# -*- coding: utf-8 -*-
"""
@File    : write_excel.py
@Time    : 2025/5/27 15:29
@Author  : zhujiantao
@Version : 1.0
@Desc    : 将扣内容相同扣信息(json文件)写入到excel中
"""
import json
from openpyxl import Workbook
from openpyxl.styles import Alignment


def write_to_excel(json_path, output_path):
    with open(json_path, "r", encoding="utf-8") as f:
        data = json.load(f)

    wb = Workbook()
    ws = wb.active
    ws.title = "Fold Info"
    ws.append(["fold_content", "same_han", "same_volume", "fold_url", "pdf_page_url"])

    row = 2  # 从第2行写起，第一行为标题
    for item in data:
        fold_content = item.get("fold_content", "")
        same_han = item.get("same_han", "")
        same_volume = item.get("same_volume", "")
        id_list = item.get("fold_id_list", [])

        for entry in id_list:
            fold_url = entry.get("fold_url", "")
            pdf_url = entry.get("pdf_page_url", "")
            ws.cell(row=row, column=4, value=fold_url)
            ws.cell(row=row, column=5, value=pdf_url)
            row += 1

        if id_list:
            start_row = row - len(id_list)
            end_row = row - 1
            for col, val in enumerate([fold_content, same_han, same_volume], start=1):
                ws.merge_cells(start_row=start_row, start_column=col, end_row=end_row, end_column=col)
                cell = ws.cell(row=start_row, column=col, value=val)
                cell.alignment = Alignment(vertical="center", horizontal="center")

    wb.save(output_path)
    print(f"✅ 成功写入 Excel：{output_path}")


write_to_excel("result.json", "tmp_result_output.xlsx")
