import re
import openpyxl as xl
from openpyxl.cell import Cell
from openpyxl.styles import PatternFill, Font

files = '16-08-05.xlsx', 'folds-0814.xlsx'
src_wb, wb = xl.open(files[0]), xl.open(files[1])
src_ws, ws = src_wb and src_wb.active, wb.active
last_cells, hans, last_han = {}, {}, ''

for i in range(2, 6402):
    # 2函序，3千字文，4经名，5卷，6扣
    c_juan, c_fold = src_ws.cell(i, 5), src_ws.cell(i, 6)
    for c in [2, 3, 4, 5, 6]:
        if c > 5 or isinstance(src_ws.cell(i, c), Cell):
            last_cells[c] = src_ws.cell(i, c)
    han, kc, jn, juan, fold = [last_cells[c].value for c in [2, 3, 4, 5, 6]]
    if juan:
        if last_cells[2].font.strike:
            han = ''
        if re.search(r'\d{2,3}', kc or ''):
            han = han or re.search(r'\d{2,3}', kc).group()
            kc = re.sub(r'\d', '', kc)
        han = re.sub(r'[^\d]', '', str(han or ''))
        if han:
            han, fold = map(int, [han, fold or 0])
            if last_han != han:
                assert han not in hans
                hans[han] = hans.get(han, [])
            if c_juan.value is None or c_juan.font.strike:  # 同卷合并扣数
                hans[han][-1] += int(c_fold.value or 0)
            else:
                hans[han].append(fold)
            last_han = han
assert len(hans) == 548

for i in range(2, 6000):
    text, fold_n = ws.cell(i, 2).value, ws[f'E{i}'].value
    if not text:
        break
    han, vol, fold_n = map(int, text.split('_') + [fold_n])
    old_n = hans[han][vol - 1] if vol <= len(hans[han]) else 0
    if old_n:  # 6 出版扣数，7 扣数差异
        ws.cell(i, 6, old_n)
        ws.cell(i, 7, fold_n - old_n)
        if fold_n != old_n:
            ws.cell(i, 7).font = Font(color='ff0000')
    else:
        ws.cell(i, 6).fill = PatternFill('solid', '1874CD')
wb.save('思溪藏扣数统计-0815.tmp.xlsx')
